# -*- coding: utf-8 -*-
"""Formatação comum dos volumes da candidatura fictícia."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

NAVY = RGBColor(0x0B, 0x1F, 0x33)
COPPER = RGBColor(0xC4, 0x62, 0x2D)
INK = RGBColor(0x1C, 0x24, 0x2C)
MUTED = RGBColor(0x4A, 0x55, 0x60)
SAND = "F3EBDD"
NAVY_HEX = "0B1F33"
COPPER_HEX = "C4622D"
WHITE = "FFFFFF"
ROW_ALT = "F7F4EF"

DISCLAIMER = (
    "Documento fictício da simulação Eleições Presidenciais das Inteligências Artificiais — Brasil 2026. "
    "Não é candidatura real, não é propaganda eleitoral, não possui registro no TSE e não representa "
    "partido, governo, empresa ou pessoa existente. Todos os cargos, mandatos e realizações atribuídos "
    "ao candidato e à vice são inventados para o experimento."
)


def _set_run_font(run, name, size, bold=False, italic=False, color=INK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def _shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_border(cell, color="C4622D"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def _page_field(paragraph):
    run = paragraph.add_run()
    _set_run_font(run, "Liberation Sans", 9, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.header_distance = Cm(0.6)
    sec.footer_distance = Cm(0.5)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Serif"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Serif")
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.12

    for style_name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[style_name]
        st.font.name = "Liberation Sans"
        st.font.color.rgb = NAVY
        st.font.bold = True
        st.font.size = Pt(size)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Sans")
        st.paragraph_format.space_before = Pt(14 if style_name != "Heading 1" else 6)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run("CANDIDATURA FICTÍCIA  ·  ELEIÇÕES PRESIDENCIAIS DAS IAs  ·  BRASIL 2026")
    _set_run_font(r, "Liberation Sans", 8, bold=True, color=COPPER)
    hp2 = header.add_paragraph()
    r2 = hp2.add_run("Elias Monteiro Viana  ·  Aliança do Trabalho e da Regra (ATR)  ·  mandato hipotético 2027–2030")
    _set_run_font(r2, "Liberation Sans", 8, color=NAVY)
    # line under header via bottom border
    pPr = hp2._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), NAVY_HEX)
    pBdr.append(bot)
    pPr.append(pBdr)

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pPr = fp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "8")
    top.set(qn("w:space"), "1")
    top.set(qn("w:color"), COPPER_HEX)
    pBdr.append(top)
    pPr.append(pBdr)
    r = fp.add_run("Simulação política  ·  sem registro no TSE  ·  página ")
    _set_run_font(r, "Liberation Sans", 9, color=MUTED)
    _page_field(fp)

    # first page header empty-ish
    fh = sec.first_page_header
    fh.is_linked_to_previous = False
    fh.paragraphs[0].text = ""
    ff = sec.first_page_footer
    ff.is_linked_to_previous = False
    ffp = ff.paragraphs[0]
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ffp.add_run("Peça fictícia da simulação  ·  página ")
    _set_run_font(r, "Liberation Sans", 9, color=MUTED)
    _page_field(ffp)

    core = doc.core_properties
    core.author = "Elias Monteiro Viana — candidatura fictícia"
    core.category = "Simulação eleitoral fictícia"
    core.comments = "Candidatura fictícia. Sem registro no TSE. Não é propaganda real."
    core.language = "pt-BR"
    return doc


def cover(doc, volume, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("REGISTRO DE CANDIDATURA")
    _set_run_font(r, "Liberation Sans", 11, bold=True, color=COPPER)

    p = doc.add_paragraph()
    r = p.add_run("ELEIÇÕES PRESIDENCIAIS DAS IAs — BRASIL 2026")
    _set_run_font(r, "Liberation Sans", 13, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(volume)
    _set_run_font(r, "Liberation Sans", 10, bold=True, color=COPPER)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    _set_run_font(r, "Liberation Sans", 22, bold=True, color=NAVY)

    if subtitle:
        p = doc.add_paragraph()
        r = p.add_run(subtitle)
        _set_run_font(r, "Liberation Serif", 12, italic=True, color=INK)

    meta = [
        "Candidato fictício: Elias Monteiro Viana",
        "Vice fictícia: Lúcia Fernanda Okamoto",
        "Partido fictício: Aliança do Trabalho e da Regra — ATR",
        "Ideologia: desenvolvimentismo de regras",
        "Cargo simulado: Presidente da República",
        "Mandato hipotético: 1º de janeiro de 2027 a 31 de dezembro de 2030",
        "Data-base do diagnóstico: 20 de outubro de 2026",
        "Slogan: O Brasil no trilho.",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(line)
        _set_run_font(r, "Liberation Sans", 11, color=NAVY)

    callout(doc, DISCLAIMER)
    doc.add_page_break()


def h(doc, level, text):
    para = doc.add_heading(text, level=level)
    return para


def p(doc, text, *, italic=False, bold=False, size=11, space_after=8):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.12
    r = para.add_run(text)
    _set_run_font(r, "Liberation Serif", size, bold=bold, italic=italic, color=INK)
    return para


def rich(doc, runs, *, space_after=8):
    """runs: list of (text, bold, italic)"""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.12
    for text, bold, italic in runs:
        r = para.add_run(text)
        _set_run_font(r, "Liberation Serif", 11, bold=bold, italic=italic, color=INK)
    return para


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.clear()
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.left_indent = Cm(1.0)
        r = para.add_run(item)
        _set_run_font(r, "Liberation Serif", 11, color=INK)


def numbered(doc, items):
    for i, item in enumerate(items, 1):
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.75)
        para.paragraph_format.first_line_indent = Cm(-0.5)
        para.paragraph_format.space_after = Pt(3)
        r = para.add_run(f"{i}. {item}")
        _set_run_font(r, "Liberation Serif", 11, color=INK)


def callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    _shade(cell, SAND)
    cell.text = ""
    para = cell.paragraphs[0]
    r = para.add_run(text)
    _set_run_font(r, "Liberation Sans", 9, color=NAVY)
    para.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def table(doc, headers, rows, *, font=8.5):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    tbl.autofit = True
    for j, htext in enumerate(headers):
        cell = tbl.rows[0].cells[j]
        _shade(cell, NAVY_HEX)
        cell.text = ""
        para = cell.paragraphs[0]
        r = para.add_run(htext)
        _set_run_font(r, "Liberation Sans", font, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.rows[i + 1].cells[j]
            if i % 2 == 1:
                _shade(cell, ROW_ALT)
            cell.text = ""
            para = cell.paragraphs[0]
            r = para.add_run(str(val))
            _set_run_font(r, "Liberation Sans", font, color=INK)
        _prevent_row_split(tbl.rows[i + 1])
    _prevent_row_split(tbl.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl


def caption(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(8)
    r = para.add_run(text)
    _set_run_font(r, "Liberation Sans", 9, italic=True, color=MUTED)


def add_picture(doc, path, width_cm, caption_text=None):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if caption_text:
        caption(doc, caption_text)


def page_break(doc):
    doc.add_page_break()


def save(doc, path):
    doc.save(path)
    return path


def eixo(doc, numero, titulo, blocos):
    """blocos: dict com as chaves obrigatórias do programa."""
    h(doc, 1, f"Eixo {numero} — {titulo}")
    labels = [
        ("diagnostico", "1. Diagnóstico"),
        ("objetivo", "2. Objetivo político"),
        ("propostas", "3. Propostas concretas"),
        ("instrumento", "4. Instrumento de implementação"),
        ("prazo", "5. Prazo"),
        ("custo", "6. Custo ou impacto fiscal"),
        ("financiamento", "7. Fonte de financiamento"),
        ("indicador", "8. Indicador de sucesso"),
        ("riscos", "9. Riscos e dificuldades"),
        ("competencia", "Competência: o que o presidente pode e o que não pode"),
    ]
    for key, label in labels:
        h(doc, 2, label)
        val = blocos[key]
        if isinstance(val, list):
            if key == "propostas":
                for i, item in enumerate(val, 1):
                    rich(doc, [(f"Proposta {numero}.{i}. ", True, False), (item, False, False)])
            else:
                bullets(doc, val)
        else:
            p(doc, val)
